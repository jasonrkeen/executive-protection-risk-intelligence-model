"""Read-only WGI workbook/API parity. No production feed, ACLED or scoring.

Run with --live to acquire public aggregates, or --replay to verify the exact
locally retained bytes. Raw files and cell comparisons must stay out of Git.
"""
from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
import csv
import hashlib
import json
import math
import platform
from pathlib import Path
from urllib.request import urlopen

from openpyxl import load_workbook, __version__ as openpyxl_version

from src.provider_verification import (
    ROOT, WGIContract, VerificationError, unpack, verify_source, verify_indicators,
)

CONFIG_PATH = ROOT / "docs/provenance/wgi_parity_contract.json"
Key = tuple[str, int, str]


@dataclass(frozen=True)
class ParityConfig:
    workbook_url: str
    catalog_url: str
    release_label: str
    start_year: int
    end_year: int
    baseline_start: int
    baseline_end: int
    scored_dimensions: list[str]
    fetched_dimensions: list[str]
    absolute_tolerance: float
    per_page: int
    columns: dict[str, str]

    @classmethod
    def load(cls) -> ParityConfig:
        return cls(**json.loads(CONFIG_PATH.read_text(encoding="utf-8")))


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def number(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise VerificationError("Invalid numeric cell; no coercion or imputation")
    return value


def read_workbook(path: Path, config: ParityConfig, contract: WGIContract) -> tuple[dict[Key, float | None], dict[str, str]]:
    """Read only six aggregate fields, never underlying-source columns."""
    values, names = {}, {}
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if set(book.sheetnames) != {d.lower() for d in contract.dimensions}:
            raise VerificationError("Workbook sheet inventory changed")
        for sheet in book:
            # Source-mean columns begin at Q in the verified provider layout.
            # Only the aggregate/identity block A:P is inspected.
            rows = sheet.iter_rows(max_col=16, values_only=True)
            headers = next(rows)
            required = ["Economy (code)", "Economy (name)", "Year", "Governance dimension", *config.columns.values()]
            if any(headers.count(h) != 1 for h in required):
                raise VerificationError("Missing/duplicate workbook headers")
            idx = {h: headers.index(h) for h in required}
            for row in rows:
                code, name, year, dim = (row[idx[h]] for h in required[:4])
                if not isinstance(code, str) or len(code) != 3 or not isinstance(name, str):
                    raise VerificationError("Invalid workbook economy identity")
                if type(year) is not int or not config.start_year <= year <= config.end_year or dim != sheet.title:
                    raise VerificationError("Unexpected workbook year/dimension")
                if code in names and names[code] != name:
                    raise VerificationError("Ambiguous workbook economy name")
                names[code] = name
                for field, header in config.columns.items():
                    key = (code, year, f"GOV_WGI_{dim.upper()}.{field}")
                    if key in values:
                        raise VerificationError("Duplicate workbook key")
                    values[key] = number(row[idx[header]])
    finally:
        book.close()
    if not values:
        raise VerificationError("Empty workbook")
    return values, names


def country_registry(payload: dict) -> dict[str, str]:
    """Use source-3 identity metadata, not guessed ISO or fuzzy-name matches."""
    try:
        source, = payload["source"]
        concept, = source["concept"]
        rows = concept["variable"]
        if (source["id"] != "3" or concept["id"] != "country"
                or int(payload["page"]) != 1 or int(payload["pages"]) != 1
                or int(payload["total"]) != len(rows) or not rows):
            raise ValueError
        result = {r["id"]: r["value"] for r in rows}
        if len(result) != len(rows) or len(set(result.values())) != len(rows):
            raise ValueError
        if any(not isinstance(k, str) or len(k) != 3 or not isinstance(v, str) or not v for k, v in result.items()):
            raise ValueError
    except (KeyError, TypeError, ValueError):
        raise VerificationError("Invalid/incomplete source-country registry") from None
    return result


def parse_api(payload: object, indicator: str, countries: dict[str, str], updated: str,
              config: ParityConfig) -> tuple[dict[Key, float | None], set[str]]:
    meta, rows = unpack(payload)
    if str(meta.get("sourceid")) != "3" or meta.get("lastupdated") != updated:
        raise VerificationError("API source/vintage mismatch")
    inverse = {name: code for code, name in countries.items()}
    values, recovered = {}, set()
    for row in rows:
        name = row.get("country", {}).get("value")
        code = row.get("countryiso3code")
        if code == "":
            code = inverse.get(name)
            if code:
                recovered.add(code)
        if code not in countries or countries[code] != name:
            raise VerificationError("Unresolved/conflicting API economy identity")
        year = row.get("date")
        if not isinstance(year, str) or not year.isdigit() or not config.start_year <= int(year) <= config.end_year:
            raise VerificationError("Invalid API year")
        if row.get("indicator", {}).get("id") != indicator:
            raise VerificationError("Wrong API indicator")
        key = (code, int(year), indicator)
        if key in values:
            raise VerificationError("Duplicate API key")
        if "value" not in row:
            raise VerificationError("API value field absent from schema")
        values[key] = number(row["value"])
    return values, recovered


def compare(workbook: dict, api: dict, countries: dict, config: ParityConfig,
            contract: WGIContract) -> tuple[dict, list[dict], list[dict]]:
    counts, by_field, by_year = Counter(), {}, {}
    exceptions, coverage = [], []
    max_difference = 0.0
    missing_dimensions: set[str]
    for country in sorted(countries):
        for year in range(config.start_year, config.end_year + 1):
            present = Counter()
            missing_dimensions = set()
            for indicator in sorted(contract.indicator_ids):
                key = (country, year, indicator)
                w, a = workbook.get(key), api.get(key)
                if w is not None: present["workbook_non_null"] += 1
                if a is not None: present["api_non_null"] += 1
                if w is None and a is None:
                    status = "BOTH_UNAVAILABLE"
                elif w is None:
                    status = "WORKBOOK_UNAVAILABLE"
                elif a is None:
                    status = "API_UNAVAILABLE"
                else:
                    difference = abs(w - a)
                    max_difference = max(max_difference, difference)
                    tolerance = 0 if indicator.endswith(".SR") else config.absolute_tolerance
                    status = "EXACT" if w == a else "WITHIN_TOLERANCE" if difference <= tolerance else "MISMATCH"
                counts[status] += 1
                by_field.setdefault(indicator, Counter())[status] += 1
                by_year.setdefault(str(year), Counter())[status] += 1
                if status in {"EXACT", "WITHIN_TOLERANCE"}: present["matched"] += 1
                else: missing_dimensions.add(indicator.removeprefix("GOV_WGI_").split(".")[0])
                if status in {"MISMATCH", "API_UNAVAILABLE", "WORKBOOK_UNAVAILABLE"}:
                    exceptions.append({"country_code": country, "year": year, "indicator": indicator,
                                       "status": status, "workbook_row_present": key in workbook,
                                       "api_row_present": key in api})
                present["workbook_absent"] += key not in workbook
                present["api_absent"] += key not in api
                present["workbook_null"] += key in workbook and w is None
                present["api_null"] += key in api and a is None
            coverage.append({"country_code": country, "year": year,
                **{k: present[k] for k in ("workbook_non_null", "api_non_null", "matched", "workbook_absent", "api_absent", "workbook_null", "api_null")},
                "complete_36_field_parity": present["matched"] == len(contract.indicator_ids)})
            coverage[-1].update({"missing_dimensions": ";".join(sorted(missing_dimensions)),
                "scored_dimensions_complete": not missing_dimensions.intersection(config.scored_dimensions),
                "fetched_dimensions_complete": not missing_dimensions.intersection(config.fetched_dimensions)})
    status = "PARITY_EXCEPTIONS" if exceptions else "PARITY_ON_AVAILABLE_CELLS"
    if not counts["EXACT"] and not counts["WITHIN_TOLERANCE"] and not exceptions:
        status = "NO_COMPARABLE_CELLS"
    return {"counts": dict(counts), "by_indicator": by_field, "by_year": by_year,
            "max_absolute_difference": max_difference,
            "status": status,
            "coverage_complete": counts["EXACT"] + counts["WITHIN_TOLERANCE"] == len(countries) * (config.end_year - config.start_year + 1) * len(contract.indicator_ids)}, coverage, exceptions


def urls(config: ParityConfig, contract: WGIContract) -> dict[str, str]:
    base = "https://api.worldbank.org/v2"
    return {"workbook.xlsx": config.workbook_url,
        "source.json": f"{base}/source/3?format=json",
        "indicators.json": f"{base}/source/3/indicator?format=json&per_page=1000",
        "countries.json": f"{base}/sources/3/country?format=json&per_page=1000",
        **{f"{i}.json": f"{base}/country/all/indicator/{i}?source=3&date={config.start_year}:{config.end_year}&format=json&per_page={config.per_page}&page=1" for i in sorted(contract.indicator_ids)}}


def acquire(folder: Path, config: ParityConfig, contract: WGIContract) -> None:
    folder.mkdir(parents=True, exist_ok=False)
    records = []
    for filename, url in urls(config, contract).items():
        with urlopen(url, timeout=60) as response:
            body, final_url, status = response.read(), response.url, response.status
        path = folder / filename
        path.write_bytes(body)
        records.append({"file": filename, "url": url, "final_url": final_url, "http_status": status,
                        "sha256": digest(path), "bytes": len(body),
                        "retrieved_at_utc": datetime.now(timezone.utc).isoformat()})
    (folder / "manifest.json").write_text(json.dumps(records, indent=2) + "\n", encoding="utf-8")


def audit(folder: Path, output: Path, config: ParityConfig, contract: WGIContract) -> dict:
    manifest = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
    expected = urls(config, contract)
    if len(manifest) != len(expected) or {r["file"] for r in manifest} != set(expected):
        raise VerificationError("Incomplete/duplicate acquisition manifest")
    for record in manifest:
        if record["url"] != expected[record["file"]] or digest(folder / record["file"]) != record["sha256"]:
            raise VerificationError("Request or byte hash mismatch")
    def read(name: str) -> object:
        return json.loads((folder / name).read_text(encoding="utf-8"))
    source = verify_source(read("source.json"), contract)
    verify_indicators(read("indicators.json"), contract)
    countries = country_registry(read("countries.json"))
    workbook, names = read_workbook(folder / "workbook.xlsx", config, contract)
    if set(names) != set(countries):
        raise VerificationError("Workbook/API country universes differ")
    api, recovered = {}, set()
    for indicator in sorted(contract.indicator_ids):
        values, recovered_codes = parse_api(read(f"{indicator}.json"), indicator, countries, source["lastupdated"], config)
        api.update(values)
        recovered.update(recovered_codes)
    result, coverage, exceptions = compare(workbook, api, countries, config, contract)
    result.update({"data_class": "REAL_PUBLIC_AGGREGATE_COMPARISON_METADATA",
        "release_label": config.release_label, "catalog_url": config.catalog_url,
        "source": source, "country_count": len(countries), "years": [config.start_year, config.end_year],
        "absolute_tolerance": config.absolute_tolerance, "source_counts_tolerance": 0,
        "blank_api_codes_resolved_by_source_registry": [{"code": c, "name": countries[c]} for c in sorted(recovered)],
        "workbook_api_name_differences": [{"code": c, "workbook_name": names[c], "api_name": countries[c]} for c in sorted(countries) if names[c] != countries[c]],
        "exceptions": exceptions, "acled": "DOCUMENTATION_ONLY_NO_REQUESTS", "operational_readiness": "NOT_ESTABLISHED",
        "requests": manifest, "config_sha256": digest(CONFIG_PATH), "verifier_sha256": digest(Path(__file__)),
        "indicator_contract_sha256": digest(ROOT / "docs/provenance/wgi_api_contract.json"),
        "provider_verifier_sha256": digest(ROOT / "src/provider_verification.py"),
        "baseline_config_sha256": digest(ROOT / "src/config.py")})
    result["environment"] = {"python": platform.python_version(), "openpyxl": openpyxl_version}
    result["windows"] = {}
    for label, start, end in [("baseline", config.baseline_start, config.baseline_end), ("latest", config.end_year, config.end_year)]:
        subset = [r for r in coverage if start <= r["year"] <= end]
        result["windows"][label] = {"years": [start, end], "country_years": len(subset),
            "complete_country_years": sum(r["complete_36_field_parity"] for r in subset),
            "scored_dimensions_complete_country_years": sum(r["scored_dimensions_complete"] for r in subset),
            "fetched_dimensions_complete_country_years": sum(r["fetched_dimensions_complete"] for r in subset),
            "incomplete": [r for r in subset if not r["complete_36_field_parity"]]}
    output.mkdir(parents=True, exist_ok=False)
    with (output / "country_year_coverage.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(coverage[0]))
        writer.writeheader(); writer.writerows(coverage)
    summaries = []
    for code in sorted(countries):
        rows = [r for r in coverage if r["country_code"] == code]
        baseline = next(r for r in rows if r["year"] == config.baseline_end)
        latest = next(r for r in rows if r["year"] == config.end_year)
        summaries.append({"country_code": code, "provider_name": countries[code],
            "complete_years_in_calendar_window": sum(r["complete_36_field_parity"] for r in rows),
            "baseline_end_year": config.baseline_end, "baseline_end_matched_fields": baseline["matched"],
            "baseline_end_missing_dimensions": baseline["missing_dimensions"],
            "baseline_end_scored_dimensions_complete": baseline["scored_dimensions_complete"],
            "latest_year": config.end_year, "latest_matched_fields": latest["matched"],
            "latest_missing_dimensions": latest["missing_dimensions"],
            "latest_scored_dimensions_complete": latest["scored_dimensions_complete"]})
    with (output / "country_summary.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(summaries[0]))
        writer.writeheader(); writer.writerows(summaries)
    result["coverage_csv_sha256"] = digest(output / "country_year_coverage.csv")
    result["country_summary_sha256"] = digest(output / "country_summary.csv")
    (output / "receipt.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--live", action="store_true")
    mode.add_argument("--replay", action="store_true")
    parser.add_argument("--raw-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    config, contract = ParityConfig.load(), WGIContract.load()
    if args.live:
        acquire(args.raw_dir, config, contract)
    result = audit(args.raw_dir, args.output_dir, config, contract)
    print(json.dumps({k: result[k] for k in ("status", "counts", "coverage_complete")}, indent=2))


if __name__ == "__main__":
    main()
