"""Invented workbook/API fixtures, never real-provider numeric evidence."""
from dataclasses import replace
import json
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
import pytest

from src.provider_verification import WGIContract, VerificationError
from src.wgi_parity import ParityConfig, read_workbook, country_registry, parse_api, compare, number, audit, urls, digest

CONFIG = replace(ParityConfig.load(), start_year=2024, end_year=2025, baseline_start=2024, baseline_end=2024)
CONTRACT = replace(WGIContract.load(), dimensions={"PV": "Political Stability"})
IND = "GOV_WGI_PV.EST"


def workbook_fixture(path, mutation=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "pv"
    sheet.append(["Economy (code)", "Economy (name)", "Year", "Governance dimension", *CONFIG.columns.values()])
    sheet.append(["AAA", "Fictional A", 2024, "pv", 0, .2, 5, 50, 40, 60])
    if mutation:
        mutation(book, sheet)
    book.save(path)


def api_payload(value=0, **changes):
    row = {"countryiso3code": "AAA", "country": {"value": "Fictional A"},
           "date": "2024", "indicator": {"id": IND}, "value": value}
    row.update(changes)
    return [{"page": 1, "pages": 1, "total": 1, "sourceid": "3", "lastupdated": "2026-09-25"}, [row]]


def registry_payload():
    return {"page": 1, "pages": 1, "total": 1, "source": [{"id": "3", "concept": [
        {"id": "country", "variable": [{"id": "AAA", "value": "Fictional A"}]}]}]}


def test_workbook_maps_units_and_preserves_zero(tmp_path):
    path = tmp_path / "toy.xlsx"
    workbook_fixture(path)
    values, names = read_workbook(path, CONFIG, CONTRACT)
    assert values[("AAA", 2024, IND)] == 0
    assert values[("AAA", 2024, "GOV_WGI_PV.SE")] == .2
    assert values[("AAA", 2024, "GOV_WGI_PV.SC_LB")] == 40
    assert len(values) == 6 and names == {"AAA": "Fictional A"}


@pytest.mark.parametrize("change", ["sheet", "header", "duplicate", "formula", "year", "dimension", "nonfinite"])
def test_workbook_schema_drift_rejected(tmp_path, change):
    def mutate(book, sheet):
        if change == "sheet": book.create_sheet("unexpected")
        if change == "header": sheet.cell(1, 5, "percentile rank")
        if change == "duplicate": sheet.append([c.value for c in sheet[2]])
        if change == "formula": sheet.cell(2, 5, "=1+1")
        if change == "year": sheet.cell(2, 3, 2026)
        if change == "dimension": sheet.cell(2, 4, "va")
        if change == "nonfinite": sheet.cell(2, 5, "NaN")
    path = tmp_path / "toy.xlsx"
    workbook_fixture(path, mutate)
    with pytest.raises(VerificationError): read_workbook(path, CONFIG, CONTRACT)


def test_blank_codes_resolved_only_from_unique_provider_registry():
    countries = country_registry(registry_payload())
    values, recovered = parse_api(api_payload(countryiso3code=""), IND, countries, "2026-09-25", CONFIG)
    assert ("AAA", 2024, IND) in values and recovered == {"AAA"}
    with pytest.raises(VerificationError):
        parse_api(api_payload(countryiso3code="", country={"value": "Fictional A typo"}), IND, countries, "2026-09-25", CONFIG)


@pytest.mark.parametrize("change", ["duplicate_name", "duplicate_code", "pages", "source", "total"])
def test_registry_ambiguity_or_incompleteness_rejected(change):
    p = registry_payload()
    if change in {"duplicate_name", "duplicate_code"}:
        p["source"][0]["concept"][0]["variable"].append({"id": "BBB" if change == "duplicate_name" else "AAA", "value": "Fictional A" if change == "duplicate_name" else "Other"})
        p["total"] = 2
    if change == "pages": p["pages"] = 2
    if change == "source": p["source"][0]["id"] = "75"
    if change == "total": p["total"] = 10
    with pytest.raises(VerificationError): country_registry(p)


@pytest.mark.parametrize("change", ["name", "code", "year", "indicator", "duplicate", "pages", "vintage", "source", "schema"])
def test_api_identity_and_completeness_fail_closed(change):
    p = api_payload()
    if change == "name": p[1][0]["country"]["value"] = "Other"
    if change == "code": p[1][0]["countryiso3code"] = "BBB"
    if change == "year": p[1][0]["date"] = "2026"
    if change == "indicator": p[1][0]["indicator"]["id"] = "PV.EST"
    if change == "duplicate": p[1].append(p[1][0]); p[0]["total"] = 2
    if change == "pages": p[0]["pages"] = 2
    if change == "vintage": p[0]["lastupdated"] = "2025-09-25"
    if change == "source": p[0]["sourceid"] = "75"
    if change == "schema": del p[1][0]["value"]
    with pytest.raises(VerificationError): parse_api(p, IND, {"AAA": "Fictional A"}, "2026-09-25", CONFIG)


@pytest.mark.parametrize("value", [True, float("nan"), float("inf"), "1.2", "#N/A"])
def test_bad_numeric_values_rejected(value):
    with pytest.raises(VerificationError): number(value)


def test_missing_cells_are_not_zero_or_successful_parity():
    key = ("AAA", 2024, IND)
    result, coverage, exceptions = compare({key: 0}, {key: None}, {"AAA": "Fictional A"}, CONFIG, CONTRACT)
    assert result["counts"]["API_UNAVAILABLE"] == 1
    assert result["counts"]["BOTH_UNAVAILABLE"] == 11
    assert not result["coverage_complete"]
    assert not coverage[0]["complete_36_field_parity"]
    assert coverage[0]["api_null"] == 1 and coverage[0]["api_absent"] == 5
    assert exceptions[0]["api_row_present"] is True


def test_precision_tolerance_is_absolute_and_counts_are_exact():
    keys = [("AAA", 2024, i) for i in sorted(CONTRACT.indicator_ids)]
    w = dict.fromkeys(keys, 1.0); a = dict(w)
    a[("AAA", 2024, IND)] += 0.5e-7
    a[("AAA", 2024, "GOV_WGI_PV.SR")] += 0.5e-7
    result, _, exceptions = compare(w, a, {"AAA": "Fictional A"}, CONFIG, CONTRACT)
    assert result["counts"]["WITHIN_TOLERANCE"] == 1
    assert len(exceptions) == 1 and exceptions[0]["indicator"].endswith(".SR")
    a[("AAA", 2024, IND)] = 1.00001
    result, _, _ = compare(w, a, {"AAA": "Fictional A"}, CONFIG, CONTRACT)
    assert result["counts"]["MISMATCH"] == 2


def test_full_parity_still_reports_unavailable_calendar_years():
    keys = [("AAA", 2024, i) for i in CONTRACT.indicator_ids]
    result, coverage, _ = compare(dict.fromkeys(keys, 0), dict.fromkeys(keys, 0), {"AAA": "A"}, CONFIG, CONTRACT)
    assert result["status"] == "PARITY_ON_AVAILABLE_CELLS"
    assert not result["coverage_complete"]
    assert coverage[0]["complete_36_field_parity"] and not coverage[1]["complete_36_field_parity"]


def test_replay_rejects_tampering_before_parsing(tmp_path):
    manifest = []
    for name, url in urls(CONFIG, CONTRACT).items():
        (tmp_path / name).write_text("not a real provider response")
        manifest.append({"file": name, "url": url, "sha256": digest(tmp_path / name)})
    (tmp_path / "manifest.json").write_text(json.dumps(manifest))
    (tmp_path / "source.json").write_text("changed")
    with patch("src.wgi_parity.urlopen", side_effect=AssertionError("replay must be offline")):
        with pytest.raises(VerificationError, match="hash mismatch"):
            audit(tmp_path, tmp_path / "result", CONFIG, CONTRACT)
    assert not (tmp_path / "result").exists()


def test_scored_and_informational_dimension_coverage_are_separate():
    c = WGIContract.load()
    keys = [("AAA", 2024, i) for i in c.indicator_ids if "_RQ." not in i]
    _, coverage, _ = compare(dict.fromkeys(keys, 0), dict.fromkeys(keys, 0), {"AAA": "A"}, CONFIG, c)
    assert coverage[0]["scored_dimensions_complete"] and coverage[0]["fetched_dimensions_complete"]
    assert not coverage[0]["complete_36_field_parity"]
    assert coverage[0]["missing_dimensions"] == "RQ"


def test_verification_scope_matches_unchanged_baseline_config():
    from src.config import START_YEAR, END_YEAR, WGI_INDICATORS
    c = ParityConfig.load()
    assert (c.baseline_start, c.baseline_end) == (START_YEAR, END_YEAR)
    assert {x.split('.')[0] for x in WGI_INDICATORS} == set(c.fetched_dimensions)


def test_no_comparable_values_cannot_claim_parity():
    result, _, _ = compare({}, {}, {"AAA": "A"}, CONFIG, CONTRACT)
    assert result["status"] == "NO_COMPARABLE_CELLS"
    assert not result["coverage_complete"]


def test_end_to_end_offline_replay_is_deterministic_and_metadata_only(tmp_path):
    raw = tmp_path / "raw"
    raw.mkdir()
    workbook_fixture(raw / "workbook.xlsx")
    source = [{"page": 1, "pages": 1, "total": 1}, [{"id": "3", "name": CONTRACT.source_name, "lastupdated": "2026-09-25"}]]
    indicators = [{"page": 1, "pages": 1, "total": 6}, [
        {"id": f"GOV_WGI_PV.{f}", "name": f"Political Stability - {n}", "source": {"id": "3"}}
        for f, n in CONTRACT.fields.items()]]
    for name, p in [("source.json", source), ("countries.json", registry_payload()), ("indicators.json", indicators)]:
        (raw / name).write_text(json.dumps(p))
    vals = {"EST": 0, "SE": .2, "SR": 5, "SC": 50, "SC_LB": 40, "SC_UB": 60}
    for field, value in vals.items():
        indicator = f"GOV_WGI_PV.{field}"
        p = api_payload(value, indicator={"id": indicator})
        (raw / f"{indicator}.json").write_text(json.dumps(p))
    manifest = [{"file": n, "url": u, "sha256": digest(raw / n)} for n, u in urls(CONFIG, CONTRACT).items()]
    (raw / "manifest.json").write_text(json.dumps(manifest))
    with patch("src.wgi_parity.urlopen", side_effect=AssertionError("offline replay")):
        first = audit(raw, tmp_path / "first", CONFIG, CONTRACT)
        second = audit(raw, tmp_path / "second", CONFIG, CONTRACT)
    assert first == second
    assert first["counts"]["EXACT"] == 6
    assert not first["coverage_complete"]
    assert first["exceptions"] == []
    assert "countryiso3code" not in json.dumps(first)
    assert (tmp_path / "first/country_year_coverage.csv").read_bytes() == (tmp_path / "second/country_year_coverage.csv").read_bytes()
    with pytest.raises(FileExistsError): audit(raw, tmp_path / "first", CONFIG, CONTRACT)
